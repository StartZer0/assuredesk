"""Authenticated API and small server-rendered review interface."""

from datetime import date
import hashlib
import json
from pathlib import Path
import secrets
import time

from argon2 import PasswordHasher
from argon2.exceptions import VerificationError
from fastapi import FastAPI, Request
from fastapi.responses import HTMLResponse, JSONResponse, Response
from fastapi.templating import Jinja2Templates
from starlette.concurrency import run_in_threadpool
from sqlalchemy import select

from . import policy, service
from .config import Settings
from pydantic import ValidationError

from .db import (
    Approval,
    Event,
    QuestionVersion,
    Chunk,
    Database,
    Draft,
    Job,
    Login,
    Objects,
    Question,
    Run,
    Source,
    User,
    Workbook,
)
from .parsing import parse_source
from .schema import DraftInput
from .workbooks import sample_workbook

PASSWORDS = PasswordHasher()


def hash_token(token):
    return hashlib.sha256(token.encode()).hexdigest()


def create_app(settings=None, db=None, objects=None):
    settings = settings or Settings.env()
    db = db or Database(settings.database_url)
    objects = objects or Objects(settings.object_root)
    templates = Jinja2Templates(directory=str(Path(__file__).with_name("templates")))
    app = FastAPI(title="AssureDesk")
    app.state.db, app.state.objects = db, objects
    login_attempts = {}

    @app.exception_handler(policy.Rejected)
    async def rejected(request, exc):
        return JSONResponse({"error": exc.code}, status_code=exc.status)

    @app.exception_handler(ValidationError)
    @app.exception_handler(ValueError)
    @app.exception_handler(KeyError)
    async def malformed(request, exc):
        return JSONResponse({"error": "invalid_request"}, status_code=422)

    @app.middleware("http")
    async def boundaries(request, call_next):
        if request.method not in {"GET", "HEAD", "OPTIONS"}:
            origin = request.headers.get("origin")
            policy_origin = str(request.base_url).rstrip("/")
            if origin and origin != policy_origin:
                return JSONResponse({"error": "origin_denied"}, status_code=403)
            length = request.headers.get("content-length")
            if length and (not length.isdigit() or int(length) > 11 * 1024**2):
                return JSONResponse({"error": "upload_too_large"}, status_code=413)
        if request.method not in {"GET", "HEAD", "OPTIONS"}:
            data = bytearray()
            async for part in request.stream():
                data.extend(part)
                if len(data) > 11 * 1024**2:
                    return JSONResponse({"error": "upload_too_large"}, status_code=413)
            request._body = bytes(data)
        response = await call_next(request)
        response.headers["Cache-Control"] = "no-store"
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["Content-Security-Policy"] = (
            "default-src 'self'; script-src 'self'; style-src 'self'; frame-ancestors 'none'; form-action 'self'"
        )
        response.headers["Referrer-Policy"] = "no-referrer"
        return response

    def principal(s, request, mutate=False):
        token = request.cookies.get("assuredesk", "")
        login = s.scalar(select(Login).where(Login.token_hash == hash_token(token)))
        policy.require(login and login.expires > time.time(), "authentication_required", 401)
        user = s.get(User, login.user_id)
        policy.require(user and user.active, "authentication_required", 401)
        if mutate:
            policy.require(
                secrets.compare_digest(request.headers.get("x-csrf-token", ""), login.csrf),
                "csrf_denied",
                403,
            )
        return user, login

    @app.get("/health/live")
    def live():
        return {"status": "alive"}

    @app.get("/health/ready")
    def ready():
        try:
            with db.transaction() as s:
                policy.require(policy.epoch(s), "schema_unavailable", 503)
                if db.engine.dialect.name == "postgresql":
                    from sqlalchemy import text

                    policy.require(
                        s.scalar(text("SELECT version_num FROM alembic_version")) == "003",
                        "migration_required",
                        503,
                    )
            return {"status": "ready"}
        except Exception:
            return JSONResponse({"error": "database_unavailable"}, status_code=503)

    @app.get("/", response_class=HTMLResponse)
    def home(request: Request):
        return templates.TemplateResponse(request=request, name="index.html", context={})

    @app.get("/app.js")
    def script():
        return Response(
            Path(__file__).with_name("templates").joinpath("app.js").read_text(),
            media_type="application/javascript",
        )

    @app.get("/style.css")
    def style():
        return Response(
            Path(__file__).with_name("templates").joinpath("style.css").read_text(), media_type="text/css"
        )

    @app.post("/v1/login")
    async def login(request: Request):
        key = request.client.host if request.client else "local"
        now = time.time()
        recent = [x for x in login_attempts.get(key, []) if x > now - 60]
        policy.require(len(recent) < 10, "login_rate_limit", 429)
        login_attempts[key] = recent + [now]
        payload = await request.json()
        with db.transaction() as s:
            user = s.scalar(select(User).where(User.name == payload.get("name")))
            try:
                valid = user and user.active and PASSWORDS.verify(user.password, payload.get("password", ""))
            except VerificationError:
                valid = False
            policy.require(valid, "invalid_login", 401)
            token, csrf = secrets.token_urlsafe(32), secrets.token_urlsafe(32)
            s.add(Login(token_hash=hash_token(token), csrf=csrf, user_id=user.id, expires=now + 8 * 3600))
        response = JSONResponse({"csrf": csrf})
        response.set_cookie(
            "assuredesk",
            token,
            secure=settings.secure_cookie,
            httponly=True,
            samesite="lax",
            max_age=8 * 3600,
        )
        return response

    @app.post("/v1/logout")
    def logout(request: Request):
        with db.transaction() as s:
            _, session = principal(s, request, True)
            s.delete(session)
        response = JSONResponse({"ok": True})
        response.delete_cookie("assuredesk")
        return response

    @app.get("/v1/me")
    def me(request: Request):
        with db.transaction() as s:
            user, login = principal(s, request)
            return {"name": user.name, "roles": user.roles, "accounts": user.accounts, "csrf": login.csrf}

    @app.get("/v1/template")
    def template(request: Request):
        with db.transaction() as s:
            principal(s, request)
        return Response(
            sample_workbook(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="questionnaire.xlsx"'},
        )

    @app.get("/v1/questionnaires")
    def list_books(request: Request):
        with db.transaction() as s:
            user, _ = principal(s, request)
            return [
                {"id": b.id, "account": b.account, "as_of": b.as_of, "confirmed": b.confirmed}
                for b in s.scalars(select(Workbook))
                if policy.content_user(user, b.account)
            ]

    @app.post("/v1/questionnaires/import", status_code=201)
    async def import_book(request: Request):
        with db.transaction() as s:
            principal(s, request, True)
        form = await request.form(max_files=1, max_fields=8, max_part_size=5 * 1024**2)
        upload = form["file"]
        data = await upload.read(5 * 1024**2 + 1)
        mapping = json.loads(str(form.get("mapping", "{}")))
        account, as_of = str(form["account"]), str(form.get("as_of", date.today().isoformat()))
        rows = await run_in_threadpool(parse_source, data, ".xlsx", mapping)
        with db.transaction() as s:
            user, _ = principal(s, request, True)
            return service.once(
                s,
                user,
                "import",
                request.headers.get("idempotency-key"),
                {"hash": hash_token(data.hex()), "mapping": mapping, "account": account, "as_of": as_of},
                lambda: {
                    "id": service.import_workbook(
                        s, user, objects, data, account, as_of, mapping, parsed_rows=rows
                    ).id
                },
            )

    @app.post("/v1/questionnaires/{ident}/mapping")
    def mapping(ident: str, request: Request):
        with db.transaction() as s:
            user, _ = principal(s, request, True)
            service.confirm_mapping(s, user, ident)
        return {"ok": True}

    @app.get("/v1/questionnaires/{ident}")
    def get_book(ident: str, request: Request):
        with db.transaction() as s:
            user, _ = principal(s, request)
            book = s.get(Workbook, ident)
            policy.require(book is not None)
            policy.access(user, book)
            questions = []
            for q in s.scalars(select(Question).where(Question.workbook_id == ident).order_by(Question.row)):
                draft = s.get(Draft, q.current_draft) if q.current_draft else None
                item = {
                    "id": q.id,
                    "key": q.key,
                    "text": q.text,
                    "row": q.row,
                    "revision": q.revision,
                    "blocked": q.blocked,
                    "coordinates": {
                        key: __import__("openpyxl").utils.get_column_letter(book.mapping[key]) + str(q.row)
                        for key in ("id_col", "question_col", "answer_col", "explanation_col")
                    },
                    "draft": None,
                }
                if draft:
                    try:
                        policy.draft_access(s, user, book, draft)
                        approved = s.scalar(select(Approval).where(Approval.draft_id == draft.id))
                        current = False
                        if approved and approved.expires > time.time():
                            try:
                                current = approved.fingerprint == service.fingerprint(s, user, q, book, draft)
                            except policy.Rejected:
                                pass
                        item["draft"] = {
                            "id": draft.id,
                            "answer": draft.answer,
                            "explanation": draft.explanation,
                            "evidence": draft.evidence,
                            "supported": draft.supported,
                            "claims": draft.claims,
                            "support_review": approved.support_review if approved else None,
                            "approved": current,
                        }
                    except policy.Rejected:
                        item["restricted_review"] = True
                latest = s.scalar(
                    select(Run)
                    .outerjoin(Event, (Event.target == Run.id) & (Event.action == "run_requested"))
                    .where(Run.question_id == q.id, Run.principal_id == user.id, Run.revision == q.revision)
                    .order_by(Event.created.desc().nulls_last(), Run.id.desc())
                    .limit(1)
                )
                if latest:
                    meta = policy.epoch(s)
                    if (latest.policy, latest.corpus) == (meta.policy, meta.corpus):
                        item["run"] = {
                            "id": latest.id,
                            "status": latest.status,
                            "stop_reason": latest.stop_reason,
                            "rounds": latest.rounds,
                            "trace": latest.record.get("trace", []),
                            "requirements": latest.record.get("proposal", {}).get("requirements", []),
                            "missing": latest.record.get("support", {}).get("missing", []),
                        }
                questions.append(item)
            return {
                "id": ident,
                "account": book.account,
                "as_of": book.as_of,
                "confirmed": book.confirmed,
                "mapping": book.mapping,
                "questions": questions,
            }

    @app.post("/v1/questions/{ident}/runs", status_code=202)
    async def run(ident: str, request: Request):
        body = await request.json()
        with db.transaction() as s:
            user, _ = principal(s, request, True)
            service.question_context(s, user, ident)
            return service.once(
                s,
                user,
                "run:" + ident,
                request.headers.get("idempotency-key"),
                body,
                lambda: {
                    "id": service.start_run(
                        s, user, ident, body["expected_revision"], body.get("mode", "corrective")
                    ).id
                },
            )

    @app.get("/v1/runs/{ident}")
    def get_run(ident: str, request: Request):
        with db.transaction() as s:
            user, _ = principal(s, request)
            run = s.get(Run, ident)
            policy.require(run and run.principal_id == user.id)
            service.question_context(s, user, run.question_id)
            meta = policy.epoch(s)
            policy.require((meta.policy, meta.corpus) == (run.policy, run.corpus), "stale_run", 409)
            return {
                "id": run.id,
                "status": run.status,
                "rounds": run.rounds,
                "generation_attempts": run.generation_attempts,
                "embedding_attempts": run.embedding_attempts,
                "stop_reason": run.stop_reason,
                "trace": run.record.get("trace", []),
                "jobs": [
                    {"id": j.id, "status": j.status}
                    for j in s.scalars(select(Job).where(Job.run_id == run.id))
                ],
            }

    @app.patch("/v1/questions/{ident}")
    async def edit_question(ident: str, request: Request):
        body = await request.json()
        with db.transaction() as s:
            user, _ = principal(s, request, True)
            policy.epoch(s, lock=True)
            q, _ = service.question_context(s, user, ident)
            policy.require(q.revision == body["expected_revision"], "stale_revision", 409)
            value = body["text"]
            policy.require(isinstance(value, str) and 0 < len(value.strip()) <= 8000, "invalid_question", 422)
            service.remember_question(s, user, q)
            previous = q.revision
            q.text, q.revision, q.current_draft = value, q.revision + 1, None
            service.remember_question(s, user, q)
            service.audit(s, user, "question_edited", q.id, before=previous, after=q.revision)
            return {"revision": q.revision}

    @app.post("/v1/questions/{ident}/drafts", status_code=201)
    async def draft(ident: str, request: Request):
        body = DraftInput.model_validate(await request.json())
        with db.transaction() as s:
            user, _ = principal(s, request, True)
            d = service.new_draft(
                s,
                user,
                ident,
                body.expected_revision,
                body.answer,
                body.explanation,
                body.evidence,
                claims=body.claims,
            )
            return {"id": d.id, "revision": d.revision}

    @app.post("/v1/questions/{ident}/decisions")
    async def decision(ident: str, request: Request):
        body = await request.json()
        with db.transaction() as s:
            user, _ = principal(s, request, True)

            _, target_book = service.question_context(s, user, ident)
            if body.get("decision") != "request_changes":
                policy.security(user, target_book)

            def apply():
                if body["decision"] == "approve":
                    result = service.approve(
                        s,
                        user,
                        ident,
                        body["expected_revision"],
                        body.get("attested_support") is True,
                        claim_reviews=body.get("claim_reviews"),
                        reason=body.get("reason", "Evidence and full answer reviewed"),
                    )
                    s.flush()
                    return {"id": result.id}
                policy.require(
                    body["decision"] in {"block", "unblock", "request_changes"}, "invalid_decision", 422
                )
                q, _ = service.question_context(s, user, ident)
                policy.require(q.revision == body["expected_revision"], "stale_revision", 409)
                service.block_question(
                    s, user, ident, body["decision"] != "unblock", body.get("reason", "Review requested")
                )
                return {"ok": True}

            return service.once(
                s, user, "decision:" + ident, request.headers.get("idempotency-key"), body, apply
            )

    @app.post("/v1/questionnaires/{ident}/exports", status_code=201)
    def export(ident: str, request: Request):
        # Creation is content-idempotent; every download revalidates the complete current manifest.
        with db.transaction() as s:
            user, _ = principal(s, request, True)
            user_id = user.id
        return {"id": service.build_export(db, user_id, objects, ident)}

    @app.get("/v1/exports/{ident}/download")
    def download(ident: str, request: Request):
        with db.transaction() as s:
            user, _ = principal(s, request)
            data = service.download(s, user, objects, ident)
        return Response(
            data,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="approved-questionnaire.xlsx"'},
        )

    @app.get("/v1/evidence/{ident}")
    def evidence(ident: str, workbook_id: str, request: Request):
        with db.transaction() as s:
            user, _ = principal(s, request)
            book = s.get(Workbook, workbook_id)
            policy.require(book is not None)
            policy.access(user, book)
            chunk, source = policy.evidence(s, user, book, [ident], disclose=False)[0]
            return {
                "id": chunk.id,
                "text": chunk.text,
                "context": chunk.parent_text,
                "title": source.title,
                "locator": chunk.locator,
                "disclosure": source.disclosure,
                "expires": source.valid_until,
            }

    @app.get("/v1/sources")
    def sources(request: Request, account: str):
        with db.transaction() as s:
            user, _ = principal(s, request)
            rows = list(s.scalars(select(Source)))
            # Older pending rebuilds stored a generation hash; activation requires a source ID.
            generations = {x.generation: x.id for x in rows if x.generation != "legacy"}
            return [
                {
                    "id": x.id,
                    "title": x.title,
                    "active": x.active,
                    "disclosure": x.disclosure,
                    "generation": x.generation,
                    "pipeline": x.pipeline,
                    "superseded": x.superseded,
                    "family": x.family,
                    "rebuild_of": generations.get(x.pipeline.get("rebuild_of"), x.pipeline.get("rebuild_of")),
                    "expires": x.valid_until,
                }
                for x in rows
                if policy.readable(user, x, account)
            ]

    @app.get("/v1/sources/{ident}/chunks")
    def source_chunks(ident: str, workbook_id: str, request: Request):
        with db.transaction() as s:
            user, _ = principal(s, request)
            book, source = s.get(Workbook, workbook_id), s.get(Source, ident)
            policy.require(book is not None and source is not None)
            policy.access(user, book)
            policy.require(policy.readable(user, source, book.account))
            return [
                {"id": c.id, "text": c.parent_text, "locator": c.locator}
                for c in s.scalars(select(Chunk).where(Chunk.source_id == ident))
            ]

    @app.post("/v1/sources/{ident}/acl")
    async def source_acl(ident: str, request: Request):
        body = await request.json()
        with db.transaction() as s:
            user, _ = principal(s, request, True)
            policy.admin(user)
            meta = policy.epoch(s, lock=True)
            src = s.get(Source, ident)
            policy.require(src is not None)
            roles = body["roles"]
            policy.require(
                isinstance(roles, list) and set(roles) <= {"sales", "security_reviewer"}, "invalid_roles", 422
            )
            src.roles = roles
            meta.policy += 1
            service.audit(s, user, "source_acl_changed", ident)
        return {"ok": True}

    @app.post("/v1/sources/import", status_code=201)
    async def source_import(request: Request):
        with db.transaction() as s:
            user, _ = principal(s, request, True)
            policy.security(user)
            user_id = user.id
        form = await request.form(max_files=1, max_fields=2, max_part_size=10 * 1024**2)
        uploaded = form["file"]
        data = await uploaded.read(10 * 1024**2 + 1)
        blocks = await run_in_threadpool(parse_source, data, Path(uploaded.filename).suffix.lower())
        metadata = json.loads(str(form["metadata"]))
        with db.transaction() as s:
            user = s.get(User, user_id)
            src = service.add_source(
                s, user, objects, data, blocks, metadata, suffix=Path(uploaded.filename).suffix.lower()
            )
            return {"id": src.id}

    @app.post("/v1/sources/{ident}/reindex", status_code=202)
    async def reindex(ident: str, request: Request):
        with db.transaction() as s:
            user, _ = principal(s, request, True)
            src = s.get(Source, ident)
            policy.require(src is not None)
            policy.curate_source(user, src.account, src.disclosure, src.audience)
            user_id = user.id
            data, suffix = objects.get(src.object_hash), src.pipeline.get("format", ".md")
            metadata = {
                key: getattr(src, key)
                for key in (
                    "family",
                    "title",
                    "account",
                    "roles",
                    "disclosure",
                    "audience",
                    "valid_from",
                    "valid_until",
                    "authority",
                    "provenance",
                )
            }
        blocks = await run_in_threadpool(parse_source, data, suffix)
        with db.transaction() as s:
            source = service.add_source(
                s, s.get(User, user_id), objects, data, blocks, metadata, suffix=suffix, rebuild_of=ident
            )
            return {"id": source.id, "supersedes": ident, "status": "pending_index"}

    @app.post("/v1/sources/{ident}/activate")
    async def activate(ident: str, request: Request):
        body = await request.json()
        with db.transaction() as s:
            user, _ = principal(s, request, True)
            service.activate_source(s, user, ident, supersedes=body.get("supersedes"))
        return {"ok": True}

    @app.post("/v1/sources/{ident}/deactivate")
    def deactivate(ident: str, request: Request):
        with db.transaction() as s:
            user, _ = principal(s, request, True)
            policy.security(user)
            meta = policy.epoch(s, lock=True)
            source = s.get(Source, ident)
            policy.require(source is not None and (source.account == "*" or source.account in user.accounts))
            policy.curate_source(user, source.account, source.disclosure, source.audience)
            source.active = False
            meta.corpus += 1
            service.audit(s, user, "source_deactivated", ident)
        return {"ok": True}

    @app.post("/v1/users/{ident}/permissions")
    async def permissions(ident: str, request: Request):
        body = await request.json()
        with db.transaction() as s:
            user, _ = principal(s, request, True)
            service.change_user(s, user, ident, body["roles"], body["accounts"], body["active"])
        return {"ok": True}

    @app.post("/v1/jobs/{ident}/cancel")
    def cancel(ident: str, request: Request):
        with db.transaction() as s:
            user, _ = principal(s, request, True)
            policy.epoch(s, lock=True)
            job = s.get(Job, ident)
            policy.require(job is not None)
            run = s.get(Run, job.run_id)
            policy.require(run.principal_id == user.id)
            q, _ = service.question_context(s, user, run.question_id)
            if run.record.get("draft_id") and run.record["draft_id"] == q.current_draft:
                q.revision += 1
                q.current_draft = None
            for delivery in s.scalars(select(Job).where(Job.run_id == run.id)):
                delivery.cancel, delivery.status, delivery.lease = True, "cancelled", 0
            run.status = "cancelled"
            service.audit(s, user, "run_cancelled", run.id)
        return {"ok": True}

    @app.get("/v1/questionnaires/{ident}/preview")
    def export_preview(ident: str, request: Request):
        with db.transaction() as s:
            user, _ = principal(s, request)
            book = s.get(Workbook, ident)
            policy.require(book is not None)
            policy.access(user, book)
            fingerprint, rows = service.export_manifest(s, user, book)
            from openpyxl.utils import get_column_letter

            return {
                "fingerprint": fingerprint,
                "count": len(rows),
                "rows": [
                    {
                        "answer_cell": f"{get_column_letter(book.mapping['answer_col'])}{row}",
                        "explanation_cell": f"{get_column_letter(book.mapping['explanation_col'])}{row}",
                        "answer": answer,
                        "explanation": explanation,
                    }
                    for row, answer, explanation in rows
                ],
            }

    @app.get("/v1/questions/{ident}/evidence-options")
    def evidence_options(ident: str, request: Request):
        with db.transaction() as s:
            user, _ = principal(s, request)
            _, book = service.question_context(s, user, ident)
            permitted = {src.id: src for src in s.scalars(select(Source)) if policy.eligible(user, src, book)}
            return [
                {"id": c.id, "title": permitted[c.source_id].title, "locator": c.locator, "text": c.text}
                for c in s.scalars(select(Chunk).where(Chunk.source_id.in_(permitted)).limit(200))
            ]

    @app.get("/v1/questions/{ident}/history")
    def history(ident: str, request: Request):
        with db.transaction() as s:
            user, _ = principal(s, request)
            q, book = service.question_context(s, user, ident)
            versions = [
                {"revision": v.revision, "text": v.text}
                for v in s.scalars(
                    select(QuestionVersion)
                    .where(QuestionVersion.question_id == ident)
                    .order_by(QuestionVersion.revision)
                )
            ]
            drafts, targets = [], [ident]
            for d in s.scalars(select(Draft).where(Draft.question_id == ident).order_by(Draft.revision)):
                try:
                    policy.draft_access(s, user, book, d)
                except policy.Rejected:
                    continue
                drafts.append(
                    {
                        "revision": d.revision,
                        "answer": d.answer,
                        "explanation": d.explanation,
                        "claims": d.claims,
                        "evidence": d.evidence,
                    }
                )
                targets.append(d.id)
            events = [
                {
                    "action": e.action,
                    "created": e.created,
                    "details": e.details
                    if "security_reviewer" in user.roles
                    else {k: v for k, v in e.details.items() if k != "reason"},
                }
                for e in s.scalars(select(Event).where(Event.target.in_(targets)).order_by(Event.created))
            ]
            return {"questions": versions, "drafts": drafts, "events": events}

    @app.get("/v1/search")
    def reviewer_search(request: Request, workbook_id: str, query: str):
        with db.transaction() as s:
            user, _ = principal(s, request)
            book = s.get(Workbook, workbook_id)
            policy.require(book is not None)
            policy.security(user, book)
            policy.require(0 < len(query.strip()) <= 1000, "invalid_query", 422)
            from .retrieval import words

            wanted = set(words(query))
            sources = {
                src.id: src
                for src in s.scalars(select(Source))
                if policy.eligible(user, src, book, disclose=False)
            }
            rows = list(s.scalars(select(Chunk).where(Chunk.source_id.in_(sources))))
            ranked = sorted(rows, key=lambda c: (-len(wanted & set(words(c.text))), c.id))
            return [
                {
                    "id": c.id,
                    "title": sources[c.source_id].title,
                    "text": c.text,
                    "locator": c.locator,
                    "disclosure": sources[c.source_id].disclosure,
                }
                for c in ranked
                if wanted & set(words(c.text))
            ][:20]

    @app.get("/v1/operations")
    def operations(request: Request):
        from .operations import status

        with db.transaction() as s:
            user, _ = principal(s, request)
            policy.require(user.active and set(user.roles) & {"administrator", "security_reviewer"})
        return status(db, objects)

    @app.delete("/v1/sources/{ident}")
    def delete_source(ident: str, request: Request):
        from .operations import delete_source as purge

        with db.transaction() as s:
            user, _ = principal(s, request, True)
            user_id = user.id
        return purge(db, objects, ident, user_id)

    return app
