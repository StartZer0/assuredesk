"""Strict workbook profile. Export constructs a clean file; it never copies ZIP parts."""

from copy import copy
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

from defusedxml.ElementTree import fromstring
from defusedxml.common import DefusedXmlException
from xml.etree.ElementTree import ParseError
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from .policy import Rejected, require
from .schema import ANSWERS

ALLOWED_PARTS = {
    "[Content_Types].xml",
    "_rels/.rels",
    "docProps/app.xml",
    "docProps/core.xml",
    "xl/workbook.xml",
    "xl/_rels/workbook.xml.rels",
    "xl/styles.xml",
    "xl/theme/theme1.xml",
    "xl/sharedStrings.xml",
    "xl/worksheets/sheet1.xml",
}


def validate_package(data):
    require(len(data) <= 5 * 1024**2, "workbook_too_large", 413)
    try:
        with ZipFile(BytesIO(data)) as z:
            names = z.namelist()
            require(len(names) <= 200 and len(set(names)) == len(names), "invalid_zip", 422)
            require(set(names) <= ALLOWED_PARTS, "unsupported_workbook_part", 422)
            require(sum(i.file_size for i in z.infolist()) <= 25 * 1024**2, "zip_too_large", 413)
            for info in z.infolist():
                require(
                    not info.flag_bits & 1 and ".." not in PurePosixPath(info.filename).parts,
                    "invalid_zip",
                    422,
                )
                require(info.file_size <= max(info.compress_size, 1) * 100, "zip_expansion", 413)
                node = fromstring(z.read(info))
                for element in node.iter():
                    tag = element.tag.rsplit("}", 1)[-1]
                    require(
                        tag not in {"f", "hyperlink", "oleObject", "externalReference"},
                        "unsupported_active_content",
                        422,
                    )
                    require(element.attrib.get("TargetMode") != "External", "external_relationship", 422)
                    if tag == "Override":
                        require(
                            "macroEnabled" not in element.attrib.get("ContentType", ""), "macro_workbook", 422
                        )
    except (BadZipFile, KeyError, ValueError, DefusedXmlException, ParseError) as exc:
        raise Rejected("invalid_workbook") from exc


def load_safe(data):
    validate_package(data)
    try:
        book = load_workbook(BytesIO(data), data_only=False, keep_links=False)
    except Exception as exc:
        raise Rejected("invalid_workbook") from exc
    require(len(book.worksheets) == 1, "one_sheet_required", 422)
    sheet = book.active
    require(sheet.sheet_state == "visible", "hidden_sheet", 422)
    require(sheet.max_row <= 500 and sheet.max_column <= 30, "worksheet_bounds", 422)
    require(not any(v.hidden for v in sheet.row_dimensions.values()), "hidden_row", 422)
    require(not any(v.hidden for v in sheet.column_dimensions.values()), "hidden_column", 422)
    for row in sheet:
        for cell in row:
            require(
                cell.data_type != "f" and not cell.hyperlink and not cell.comment, "unsupported_cell", 422
            )
            require(
                cell.value is None or isinstance(cell.value, (str, int, float, bool)),
                "unsupported_cell_value",
                422,
            )
            require(len(str(cell.value if cell.value is not None else "")) <= 8000, "cell_too_large", 422)
    for validation in sheet.data_validations.dataValidation:
        require(
            validation.type == "list"
            and isinstance(validation.formula1, str)
            and validation.formula1.startswith('"')
            and validation.formula1.endswith('"'),
            "unsupported_validation",
            422,
        )
        choices = validation.formula1[1:-1].split(",")
        require(set(choices) == set(ANSWERS), "uncertainty_not_representable", 422)
    return book


def extract(data, mapping):
    book = load_safe(data)
    sheet = book.active
    keys = {"id_col", "question_col", "answer_col", "explanation_col", "first_row", "last_row"}
    require(set(mapping) == keys and all(type(v) is int for v in mapping.values()), "invalid_mapping", 422)
    cols = [mapping[k] for k in ("id_col", "question_col", "answer_col", "explanation_col")]
    require(len(set(cols)) == 4 and min(cols) >= 1 and max(cols) <= 30, "overlapping_mapping", 422)
    first, last = mapping["first_row"], mapping["last_row"]
    require(1 <= first <= last <= 500 and last - first < 50, "question_count", 422)
    require(
        not any(r.max_row >= first and r.min_row <= last for r in sheet.merged_cells.ranges),
        "merged_question_cells",
        422,
    )
    rows, seen = [], set()
    for row in range(first, last + 1):
        key = str(sheet.cell(row, mapping["id_col"]).value or "").strip()
        question = str(sheet.cell(row, mapping["question_col"]).value or "").strip()
        require(key and question and key not in seen and len(key) <= 120, "invalid_question", 422)
        seen.add(key)
        rows.append({"key": key, "text": question, "row": row})
    return rows


def clean_export(original, mapping, answers, questions=None):
    old = load_safe(original).active
    out = Workbook()
    sheet = out.active
    sheet.title = old.title
    response_cols = {mapping["answer_col"], mapping["explanation_col"]}
    for row in old:
        for cell in row:
            target = sheet.cell(cell.row, cell.column)
            is_answer = (
                mapping["first_row"] <= cell.row <= mapping["last_row"] and cell.column in response_cols
            )
            target.value = "" if is_answer else str(cell.value if cell.value is not None else "")
            target.data_type = "s"
            if cell.has_style:
                target.font = copy(cell.font)
                target.fill = copy(cell.fill)
                target.border = copy(cell.border)
                target.alignment = copy(cell.alignment)
                target.number_format = "@"
    for merged in old.merged_cells.ranges:
        sheet.merge_cells(str(merged))
    require(len(answers) == mapping["last_row"] - mapping["first_row"] + 1, "incomplete_export", 409)
    for row, answer, explanation in answers:
        require(answer in ANSWERS, "invalid_answer", 422)
        for col, text in [(mapping["answer_col"], answer), (mapping["explanation_col"], explanation)]:
            cell = sheet.cell(row, col)
            cell.value, cell.data_type = text, "s"
    for row, question in (questions or {}).items():
        cell = sheet.cell(row, mapping["question_col"])
        cell.value, cell.data_type = question, "s"
    out.properties.creator = "AssureDesk"
    out.properties.lastModifiedBy = "AssureDesk"
    stream = BytesIO()
    out.save(stream)
    validate_package(stream.getvalue())
    return stream.getvalue()


def sample_workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = "Questionnaire"
    sheet.append(["ID", "Question", "Response", "Explanation"])
    questions = [
        "Is all customer data, including backups, stored exclusively in the EU?",
        "Is customer data encrypted in transit?",
        "Is customer data encrypted at rest?",
        "How long are backups retained?",
        "Is customer data used to train general-purpose AI models?",
    ]
    for i, question in enumerate(questions, 1):
        sheet.append([f"Q{i}", question, "", ""])
    dv = DataValidation(type="list", formula1='"' + ",".join(ANSWERS) + '"')
    sheet.add_data_validation(dv)
    dv.add("C2:C6")
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
