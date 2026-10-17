from io import BytesIO
import json
from uuid import uuid4

from argon2 import PasswordHasher
from fastapi.testclient import TestClient

from app.config import Settings
from app.db import User
from app.web import create_app
from app.workbooks import sample_workbook
from conftest import MAPPING


def client(env, role="sales"):
    with env.db.transaction() as s:
        s.get(User, getattr(env, role)).password = PasswordHasher().hash("local-test-password")
    c = TestClient(create_app(Settings("sqlite://", secure_cookie=False), env.db, env.objects))
    r = c.post("/v1/login", json={"name": role, "password": "local-test-password"})
    assert r.status_code == 200
    c.headers["X-CSRF-Token"] = r.json()["csrf"]
    c.headers["Idempotency-Key"] = str(uuid4())
    return c


def test_http_auth_csrf_and_origin(env):
    c = client(env)
    assert c.get("/v1/me").json()["roles"] == ["sales"]
    assert (
        c.post(f"/v1/questionnaires/{env.book}/mapping", headers={"X-CSRF-Token": "bad"}).status_code == 403
    )
    assert c.post("/v1/logout", headers={"Origin": "https://untrusted.invalid"}).status_code == 403
    c.cookies.clear()
    assert c.get("/v1/me").status_code == 401


def test_questionnaire_displays_newest_run_by_creation_not_random_id(env, monkeypatch):
    from sqlalchemy import select
    from app import service
    from app.db import Event

    identities = iter(["ffffffff-ffff-ffff-ffff-ffffffffffff", "00000000-0000-0000-0000-000000000000"])
    monkeypatch.setattr(service, "uid", lambda: next(identities))
    with env.db.transaction() as s:
        actor = s.get(User, env.sales)
        older = service.start_run(s, actor, env.q, 1, "corrective")
        newer = service.start_run(s, actor, env.q, 1, "corrective")
        s.flush()
        for event in s.scalars(select(Event).where(Event.action == "run_requested")):
            event.created = 100 if event.target == older.id else 200
        newer_id = newer.id
    question = next(
        q for q in client(env).get(f"/v1/questionnaires/{env.book}").json()["questions"] if q["id"] == env.q
    )
    assert question["run"]["id"] == newer_id


def test_evidence_url_has_no_cross_customer_or_internal_leak(env):
    c = client(env)
    for key in ["internal", "nova"]:
        r = c.get(f"/v1/evidence/{env.chunks[key]}?workbook_id={env.book}")
        assert r.status_code == 404 and "CANARY" not in r.text
    assert c.get(f"/v1/evidence/{env.chunks['public']}?workbook_id={env.book}").status_code == 200


def test_cross_principal_run_cannot_be_hydrated(env):
    c = client(env)
    r = c.post(f"/v1/questions/{env.q}/runs", json={"expected_revision": 1}).json()
    other = client(env, "reviewer")
    assert other.get("/v1/runs/" + r["id"]).status_code == 404


def test_full_http_manual_workflow(env):
    c = client(env, "reviewer")
    r = c.post(
        "/v1/questionnaires/import",
        data={"account": "acme", "as_of": "2026-10-04", "mapping": json.dumps(MAPPING)},
        files={"file": ("test.xlsx", sample_workbook())},
    )
    assert r.status_code == 201, r.text
    ident = r.json()["id"]
    b = c.get("/v1/questionnaires/" + ident).json()
    assert not b["confirmed"] and len(b["questions"]) == 5
    assert c.post(f"/v1/questionnaires/{ident}/mapping").status_code == 200
    for q in b["questions"]:
        d = c.post(
            f"/v1/questions/{q['id']}/drafts",
            json={
                "expected_revision": 1,
                "answer": "Unknown",
                "explanation": "Not documented; reviewer approved an explicit unknown.",
                "evidence": [],
            },
        )
        assert d.status_code == 201
        c.headers["Idempotency-Key"] = str(uuid4())
        assert (
            c.post(
                f"/v1/questions/{q['id']}/decisions",
                json={"expected_revision": 2, "decision": "approve", "attested_support": True},
            ).status_code
            == 200
        )
    e = c.post(f"/v1/questionnaires/{ident}/exports")
    assert e.status_code == 201
    r = c.get(f"/v1/exports/{e.json()['id']}/download")
    assert r.status_code == 200
    from openpyxl import load_workbook

    w = load_workbook(BytesIO(r.content))
    assert w.active["C6"].value == "Unknown"


def test_admin_has_no_content_and_can_revoke(env):
    c = client(env, "admin")
    assert c.get("/v1/questionnaires").json() == []
    assert c.get(f"/v1/questionnaires/{env.book}").status_code == 404
    r = c.post(f"/v1/users/{env.sales}/permissions", json={"roles": [], "accounts": [], "active": False})
    assert r.status_code == 200


def test_invalid_json_payload_is_safe(env):
    c = client(env)
    r = c.post(f"/v1/questions/{env.q}/drafts", json={"answer": "pretend"})
    assert r.status_code == 422 and r.json() == {"error": "invalid_request"}


def test_cancel_api_reveals_only_own_jobs(env):
    c = client(env)
    r = c.post(f"/v1/questions/{env.q}/runs", json={"expected_revision": 1}).json()
    status = c.get("/v1/runs/" + r["id"]).json()
    ident = status["jobs"][0]["id"]
    assert client(env, "reviewer").post("/v1/jobs/" + ident + "/cancel").status_code == 404
    assert c.post("/v1/jobs/" + ident + "/cancel").status_code == 200
    assert c.get("/v1/runs/" + r["id"]).json()["status"] == "cancelled"


def test_question_edit_changes_export_and_invalidates_old_draft(env):
    c = client(env, "reviewer")
    q = c.patch(f"/v1/questions/{env.q}", json={"expected_revision": 1, "text": "Corrected question text?"})
    assert q.status_code == 200 and q.json()["revision"] == 2
    from test_boundaries import approve_all

    approve_all(env)
    e = c.post(f"/v1/questionnaires/{env.book}/exports").json()
    r = c.get(f"/v1/exports/{e['id']}/download")
    from openpyxl import load_workbook

    assert load_workbook(BytesIO(r.content)).active["B2"].value == "Corrected question text?"


def test_cancel_after_draft_revokes_its_exportability(env):
    from langgraph.checkpoint.memory import InMemorySaver
    from app.workflow import Engine
    from app.schema import Support
    from app.db import Question
    from app import service
    from conftest import Generator, Embeddings
    from test_boundaries import propose, user

    c = client(env)
    run = c.post(f"/v1/questions/{env.q}/runs", json={"expected_revision": 1}).json()["id"]
    Engine(
        env.db, Generator([propose, Support(supported=True, missing=[], next_queries=[])]), Embeddings()
    ).graph(InMemorySaver()).invoke({"run_id": run, "route": "search"}, {"configurable": {"thread_id": run}})
    with env.db.transaction() as s:
        for ident in env.questions:
            q = s.get(Question, ident)
            if ident != env.q:
                service.new_draft(s, user(s, env), ident, q.revision, "Unknown", "No supported answer.", [])
            service.approve(s, user(s, env, "reviewer"), ident, q.revision, True)
    artifact = service.build_export(env.db, env.sales, env.objects, env.book)
    job = c.get("/v1/runs/" + run).json()["jobs"][0]["id"]
    assert c.post("/v1/jobs/" + job + "/cancel").status_code == 200
    assert c.get("/v1/exports/" + artifact + "/download").status_code == 409
    assert all(j["status"] == "cancelled" for j in c.get("/v1/runs/" + run).json()["jobs"])
